#!/usr/bin/env python3
"""
Genera los archivos de datos de la Calculadora Logistica a partir de los
tarifarios oficiales en Excel.

Uso:
    python3 herramientas/generar_tarifas.py

Lee:
    tarifarios/usend-2026-07.xlsx   (hojas TARIFARIO y MT)
    tarifarios/starken-2025.xlsx    (hojas Starken2025 y Comunas)

Escribe:
    data/tarifas-usend.js
    data/tarifas-starken.js
    data/destinos.js

Al recibir un tarifario nuevo: reemplaza el archivo en tarifarios/, ajusta
las constantes de configuracion de abajo y vuelve a ejecutar el script.
"""

import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent

# ---------------------------------------------------------------------------
# Configuracion: actualizar al cargar un tarifario nuevo
# ---------------------------------------------------------------------------

USEND_XLSX = RAIZ / "tarifarios" / "usend-2026-07.xlsx"
USEND_HOJA_TARIFAS = "TARIFARIO"
USEND_HOJA_COMUNAS = "MT"
USEND_VIGENCIA = {
    "desde": "2026-07-11",
    "hasta": "2027-01-10",
    "etiqueta": "11 julio 2026 al 10 enero 2027",
}

STARKEN_XLSX = RAIZ / "tarifarios" / "starken-2025.xlsx"
STARKEN_HOJA_TARIFAS = "Starken2025"
STARKEN_HOJA_COMUNAS = "Comunas"
STARKEN_VIGENCIA = {
    "desde": "2025-01-01",
    "hasta": None,
    "etiqueta": "Tarifario 2025",
}

SALIDA = RAIZ / "data"

# Factor de conversion del peso volumetrico, en cm3 por kilo.
# U Send: 1 m3 = 250 kg  ->  1.000.000 / 250 = 4000
# Starken: (largo x ancho x alto) / 5000
USEND_FACTOR_VOL = 4000
STARKEN_FACTOR_VOL = 5000

# Comunas cuyo nombre difiere entre ambos tarifarios.
# clave: nombre en el tarifario U Send  ->  destino en el tarifario Starken
ALIAS_STARKEN = {
    "AYSEN": "PUERTO AYSEN",
    "NATALES": "PUERTO NATALES",
    "CISNES": "PUERTO CISNES",
    "SAN FRANCISCO MOSTAZAL": "SAN FRANCISCO DE MOSTAZAL",
    "SAN VICENTE": "SAN VICENTE DE TAGUATAGUA",
    "MARIQUINA": "SAN JOSE DE LA MARIQUINA",
    "PALMILLA": "PALMILLA SAN FERNANDO",
    "PLACILLA": "PLACILLA (SAN FERNANDO)",
    "EL CARMEN": "EL CARMEN CHILLAN",
    "OLIVAR": "OLIVAR BAJO",
    "CALERA": "LA CALERA",
    "SANTO DOMINGO": "ROCAS DE SANTO DOMINGO",
    "O'HIGGINS": "VILLA OHIGGINS",
    "ALTO BIOBIO": "ALTO BIO BIO",
    "PAIGUANO": "PAIHUANO",
    "TREGUACO": "TREHUACO",
    "SAN PEDRO": "SAN PEDRO DE MELIPILLA",
}

# Comunas sin destino propio en el tarifario Starken, referenciadas a la
# localidad mas cercana del mismo grupo tarifario. Se marcan como referenciales.
ALIAS_STARKEN_APROX = {
    "CHILLAN VIEJO": "CHILLAN",
    "HUALPEN": "HUALPENCILLO",
}

# Starken cobra todo el Gran Santiago como un unico destino.
STARKEN_DESTINO_RM = "SANTIAGO"

REGIONES = {
    15: "Arica y Parinacota",
    1: "Tarapaca",
    2: "Antofagasta",
    3: "Atacama",
    4: "Coquimbo",
    5: "Valparaiso",
    13: "Metropolitana de Santiago",
    6: "Libertador General Bernardo O'Higgins",
    7: "Maule",
    16: "Nuble",
    8: "Biobio",
    9: "La Araucania",
    14: "Los Rios",
    10: "Los Lagos",
    11: "Aysen",
    12: "Magallanes y la Antartica Chilena",
}


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def norm(valor):
    """Normaliza un nombre para comparar: mayusculas, sin tildes ni puntuacion."""
    if valor is None:
        return ""
    texto = str(valor).strip().upper().replace(".", " ")
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    return " ".join(texto.split())


def titulo(valor):
    """Nombre presentable: Puerto Montt, Vina del Mar."""
    menores = {"DE", "DEL", "LA", "LAS", "LOS", "EL", "Y"}
    palabras = norm(valor).split()
    salida = []
    for i, p in enumerate(palabras):
        salida.append(p.capitalize() if (i == 0 or p not in menores) else p.lower())
    return " ".join(salida)


def num(valor):
    """Convierte a numero, o None si la celda no es numerica."""
    if isinstance(valor, (int, float)):
        return round(float(valor), 4)
    return None


def escribir_js(ruta, variable, datos, encabezado):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    cuerpo = json.dumps(datos, ensure_ascii=False, separators=(",", ":"))
    ruta.write_text(
        "// " + encabezado + "\n"
        "// Archivo generado por herramientas/generar_tarifas.py — no editar a mano.\n"
        "window." + variable + " = " + cuerpo + ";\n",
        encoding="utf-8",
    )
    print("  escrito %-28s %7.1f KB" % (ruta.name, ruta.stat().st_size / 1024))


# ---------------------------------------------------------------------------
# U Send
# ---------------------------------------------------------------------------

TRAMOS_FIJOS = [
    {"hasta": 1, "etiqueta": "0 a 1 kg"},
    {"hasta": 2, "etiqueta": "1 a 2 kg"},
    {"hasta": 3, "etiqueta": "2 a 3 kg"},
    {"hasta": 4, "etiqueta": "3 a 4 kg"},
    {"hasta": 9, "etiqueta": "4 a 9 kg"},
    {"hasta": 15, "etiqueta": "9 a 15 kg"},
]

TRAMOS_ADICIONALES = [
    {"hasta": 50, "etiqueta": "15 a 50 kg"},
    {"hasta": 100, "etiqueta": "50 a 100 kg"},
    {"hasta": 200, "etiqueta": "100 a 200 kg"},
    {"hasta": 500, "etiqueta": "200 a 500 kg"},
    {"hasta": 1000, "etiqueta": "500 a 1000 kg"},
    {"hasta": None, "etiqueta": "sobre 1000 kg"},
]


def leer_usend(wb):
    """Devuelve (grupos, sucursales, aereo) del tarifario U Send."""
    ws = wb[USEND_HOJA_TARIFAS]
    grupos = defaultdict(dict)      # grupo -> zona -> tarifas
    sucursales = {}                 # sucursal -> grupo
    zona_actual = None
    aereo = defaultdict(dict)       # zona -> sucursal -> tarifas
    zona_aerea = None
    bloques_aereos = 0
    modo_aereo = False

    for fila in ws.iter_rows(min_row=1, max_col=20, values_only=True):
        etiqueta_a = norm(fila[1])
        if etiqueta_a.startswith("TARIFARIO AEREO"):
            modo_aereo = True
            continue

        zona_celda = norm(fila[1])
        sucursal = norm(fila[2])

        if not modo_aereo:
            m = re.match(r"^(Z[123])\s*:", zona_celda)
            if m:
                zona_actual = m.group(1)
            if not sucursal or zona_actual is None:
                continue
            grupo = num(fila[3])
            valores = [num(v) for v in fila[4:17]]
            if grupo is None or any(v is None for v in valores):
                continue
            grupo = str(int(grupo))
            sucursales.setdefault(sucursal, grupo)
            if sucursales[sucursal] != grupo:
                print("  AVISO: %s aparece con grupos distintos (%s / %s)"
                      % (sucursal, sucursales[sucursal], grupo))
            tarifas = {
                "sobre": valores[0],
                "fijos": valores[1:7],
                "adicionales": valores[7:13],
            }
            previo = grupos[grupo].get(zona_actual)
            if previo and previo != tarifas:
                print("  AVISO: grupo %s zona %s con valores distintos en %s"
                      % (grupo, zona_actual, sucursal))
            grupos[grupo][zona_actual] = tarifas
        else:
            if zona_celda.startswith("ZONA URBANA"):
                bloques_aereos += 1
                # El tarifario oficial rotula el tercer bloque como Z2 por error.
                zona_aerea = ["Z1", "Z2", "Z3"][min(bloques_aereos - 1, 2)]
            if not sucursal or zona_aerea is None:
                continue
            valores = [num(v) for v in fila[3:10] if num(v) is not None]
            if len(valores) < 3:
                continue
            aereo[zona_aerea][sucursal] = {
                "sobre": valores[0],
                "fijo": valores[1],
                "adicional": valores[2],
            }

    return dict(grupos), sucursales, {z: dict(v) for z, v in aereo.items()}


def leer_comunas_usend(wb):
    """comuna normalizada -> {agencia, zona, sla, slaAereo, provincia, region}."""
    ws = wb[USEND_HOJA_COMUNAS]
    # Columnas: 2 destino, 3 provincia, 4 region, 5 agencia, 6 zona,
    #           7 SLA terrestre, 8 SLA aereo.
    filas = {}
    for fila in ws.iter_rows(min_row=2, max_col=9, values_only=True):
        comuna = norm(fila[2])
        if not comuna:
            continue
        sla = num(fila[7])
        sla_aereo = num(fila[8])
        filas[comuna] = {
            "agencia": norm(fila[5]),
            "zona": norm(fila[6]),
            "provincia": titulo(fila[3]),
            "region": titulo(fila[4]),
            "sla": int(sla) if sla else None,
            "slaAereo": int(sla_aereo) if sla_aereo else None,
        }
    return filas


# ---------------------------------------------------------------------------
# Starken
# ---------------------------------------------------------------------------

def leer_starken(wb):
    """Devuelve (perfiles, destinos). Cada destino apunta a un perfil de precios."""
    ws = wb[STARKEN_HOJA_TARIFAS]
    encabezado = [c.value for c in ws[7]]
    col_kg = {}
    for i, valor in enumerate(encabezado):
        if isinstance(valor, (int, float)) and 16 <= valor <= 100:
            col_kg[int(valor)] = i
    col_cerrados = [2, 3, 4, 5, 6]          # 0-0,5 / 0,51-3 / 3,01-6 / 6,01-10 / 10,01-15
    col_adicional = encabezado.index(" 100 +")
    col_provincia = encabezado.index("Provincia")
    col_agencia = encabezado.index("AGENCIA DE COBERTURA")
    col_zona = encabezado.index("ZONA")
    col_region = encabezado.index("Numero Region")

    perfiles = []
    indice_perfil = {}
    destinos = {}

    for fila in ws.iter_rows(min_row=8, max_col=col_region + 1, values_only=True):
        nombre = norm(fila[1])
        if not nombre or nombre == ".":
            continue
        cerrados = [num(fila[c]) for c in col_cerrados]
        por_kilo = [num(fila[col_kg[k]]) for k in range(16, 101)]
        adicional = num(fila[col_adicional])
        if any(v is None for v in cerrados + por_kilo) or adicional is None:
            print("  AVISO: destino Starken sin precios completos: %s" % nombre)
            continue
        clave = tuple(cerrados + por_kilo + [adicional])
        if clave not in indice_perfil:
            indice_perfil[clave] = len(perfiles)
            perfiles.append({
                "cerrados": [int(round(v)) for v in cerrados],
                "porKilo": [int(round(v)) for v in por_kilo],
                "adicional": int(round(adicional)),
            })
        region = num(fila[col_region])
        destinos[nombre] = {
            "perfil": indice_perfil[clave],
            "provincia": titulo(fila[col_provincia]),
            "region": REGIONES.get(int(region)) if region else None,
            "agenciaUsend": norm(fila[col_agencia]),
            "zonaUsend": norm(fila[col_zona]),
        }
    return perfiles, destinos


# ---------------------------------------------------------------------------
# Construccion del maestro de destinos
# ---------------------------------------------------------------------------

def construir_destinos(comunas_usend, starken_destinos, sucursales_usend):
    """Une ambos tarifarios en una sola lista de destinos consultables."""
    maestro = {}

    # 1. Base: todos los destinos del tarifario Starken.
    for nombre, dato in starken_destinos.items():
        maestro[nombre] = {
            "n": titulo(nombre),
            "prov": dato["provincia"],
            "reg": dato["region"],
            "s": {"p": dato["perfil"]},
            "u": None,
        }
        agencia, zona = dato["agenciaUsend"], dato["zonaUsend"]
        if agencia in sucursales_usend and zona in ("Z1", "Z2", "Z3"):
            maestro[nombre]["u"] = {"ag": agencia, "z": zona}

    # 2. Datos de cobertura U Send: zona oficial y SLA por comuna.
    def aplicar_usend(entrada, dato):
        """Escribe la cobertura U Send, que manda sobre lo que trae la hoja de Starken."""
        if dato["agencia"] in sucursales_usend and dato["zona"] in ("Z1", "Z2", "Z3"):
            entrada["u"] = {"ag": dato["agencia"], "z": dato["zona"]}
            if dato["sla"]:
                entrada["u"]["sla"] = dato["sla"]
            if dato["slaAereo"]:
                entrada["u"]["slaA"] = dato["slaAereo"]
        else:
            entrada["u"] = None
            entrada["fc"] = True   # fuera de cobertura U Send

    sin_starken = []
    for comuna, dato in comunas_usend.items():
        destino_starken = None
        nota = None
        if comuna in starken_destinos:
            destino_starken = comuna
        elif comuna in ALIAS_STARKEN:
            destino_starken = norm(ALIAS_STARKEN[comuna])
            nota = "alias"
        elif comuna in ALIAS_STARKEN_APROX:
            destino_starken = norm(ALIAS_STARKEN_APROX[comuna])
            nota = "aprox"
        elif norm(dato["region"]) == "METROPOLITANA DE SANTIAGO":
            destino_starken = STARKEN_DESTINO_RM
            nota = "ref"

        # Mismo lugar con dos nombres: se fusiona en un solo destino consultable
        # por cualquiera de los dos. No aplica al Gran Santiago, donde cada
        # comuna mantiene su propia zona U Send.
        if nota == "alias" and destino_starken in maestro:
            entrada = maestro[destino_starken]
            alternativos = entrada.setdefault("alt", [])
            if titulo(comuna) not in alternativos:
                alternativos.append(titulo(comuna))
            entrada["prov"] = dato["provincia"] or entrada["prov"]
            entrada["reg"] = dato["region"] or entrada["reg"]
            aplicar_usend(entrada, dato)
            continue

        entrada = maestro.get(comuna)
        if entrada is None:
            entrada = {
                "n": titulo(comuna),
                "prov": dato["provincia"],
                "reg": dato["region"],
                "s": None,
                "u": None,
            }
            maestro[comuna] = entrada
        entrada["prov"] = entrada["prov"] or dato["provincia"]
        entrada["reg"] = entrada["reg"] or dato["region"]

        aplicar_usend(entrada, dato)

        if entrada["s"] is None and destino_starken in starken_destinos:
            entrada["s"] = {"p": starken_destinos[destino_starken]["perfil"]}
            if nota:
                entrada["s"]["via"] = titulo(destino_starken)
                entrada["s"]["nota"] = nota
        if entrada["s"] is None:
            sin_starken.append(comuna)

    # 3. SLA referencial para localidades sin dato propio: moda por agencia y zona.
    referencia_sla = defaultdict(Counter)
    for dato in comunas_usend.values():
        if dato["sla"] and dato["agencia"] and dato["zona"]:
            referencia_sla[(dato["agencia"], dato["zona"])][dato["sla"]] += 1
    for entrada in maestro.values():
        u = entrada.get("u")
        if u and "sla" not in u:
            comun = referencia_sla.get((u["ag"], u["z"]))
            if comun:
                u["sla"] = comun.most_common(1)[0][0]
                u["slaRef"] = True

    lista = sorted(maestro.values(), key=lambda d: d["n"])
    return lista, sin_starken


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    for ruta in (USEND_XLSX, STARKEN_XLSX):
        if not ruta.exists():
            sys.exit("No se encuentra el tarifario: %s" % ruta)

    print("Leyendo %s" % USEND_XLSX.name)
    wb_usend = openpyxl.load_workbook(USEND_XLSX, data_only=True)
    grupos, sucursales, aereo = leer_usend(wb_usend)
    comunas_usend = leer_comunas_usend(wb_usend)
    print("  %d grupos de cobro, %d sucursales, %d comunas"
          % (len(grupos), len(sucursales), len(comunas_usend)))

    print("Leyendo %s" % STARKEN_XLSX.name)
    wb_starken = openpyxl.load_workbook(STARKEN_XLSX, data_only=True)
    perfiles, starken_destinos = leer_starken(wb_starken)
    print("  %d destinos, %d perfiles de precio distintos"
          % (len(starken_destinos), len(perfiles)))

    destinos, sin_starken = construir_destinos(comunas_usend, starken_destinos, sucursales)
    con_ambos = sum(1 for d in destinos if d["u"] and d["s"])
    print("Maestro: %d destinos, %d con los dos couriers" % (len(destinos), con_ambos))
    if sin_starken:
        print("  Comunas U Send sin tarifa Starken (%d): %s"
              % (len(sin_starken), ", ".join(sorted(sin_starken))))

    print("Generando archivos de datos")
    escribir_js(
        SALIDA / "tarifas-usend.js", "TARIFAS_USEND",
        {
            "courier": "U Send",
            "vigencia": USEND_VIGENCIA,
            "incluyeIva": False,
            "factorVolumetrico": USEND_FACTOR_VOL,
            "tramosFijos": TRAMOS_FIJOS,
            "tramosAdicionales": TRAMOS_ADICIONALES,
            "grupos": grupos,
            "sucursales": sucursales,
            "aereo": aereo,
            "notas": [
                "Valores netos, no incluyen IVA.",
                "Servicio puerta a puerta, origen Santiago, sujeto a restricciones de cobertura y volumen.",
                "Se cobra la mayor relacion entre peso fisico y peso volumetrico (1 m3 = 250 kg).",
                "Seguro incluido hasta 7 UF por bulto y 60 UF por guia. Cobertura adicional 0,35% del valor declarado, 1,3% en electronica.",
                "Reajuste trimestral segun IPC.",
                "Devolucion de documentos: $700 + IVA.",
            ],
        },
        "Tarifario U Send vigente %s" % USEND_VIGENCIA["etiqueta"],
    )
    escribir_js(
        SALIDA / "tarifas-starken.js", "TARIFAS_STARKEN",
        {
            "courier": "Starken",
            "servicio": "Normal, origen Santiago",
            "vigencia": STARKEN_VIGENCIA,
            "incluyeIva": False,
            "factorVolumetrico": STARKEN_FACTOR_VOL,
            "tramosCerrados": [
                {"hasta": 0.5, "etiqueta": "0 a 0,5 kg"},
                {"hasta": 3, "etiqueta": "0,51 a 3 kg"},
                {"hasta": 6, "etiqueta": "3,01 a 6 kg"},
                {"hasta": 10, "etiqueta": "6,01 a 10 kg"},
                {"hasta": 15, "etiqueta": "10,01 a 15 kg"},
            ],
            "perfiles": perfiles,
            "notas": [
                "Valores netos, no incluyen IVA.",
                "Tarifa servicio normal, origen Santiago.",
                "Entre 16 y 100 kg el tarifario fija un valor por kilo exacto.",
                "Sobre 100 kg se suma el valor por kilo adicional del destino.",
                "Peso volumetrico = (largo x ancho x alto) / 5000.",
            ],
        },
        "Tarifario Starken %s" % STARKEN_VIGENCIA["etiqueta"],
    )
    escribir_js(
        SALIDA / "destinos.js", "DESTINOS", destinos,
        "Maestro de destinos: zona y SLA U Send + perfil tarifario Starken",
    )
    print("Listo.")


if __name__ == "__main__":
    main()
